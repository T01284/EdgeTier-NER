"""Regenerate Elsevier Highlights + Graphical Abstract upload artifacts."""
from __future__ import annotations

import os
import shutil
import subprocess
from pathlib import Path

from docx import Document
from PIL import Image

ROOT = Path(r"d:\paper\EdgeFS_NER")
SUB = ROOT / "paper" / "submission"
GA_DEL = ROOT / "paper" / "delivery" / "attachments" / "graphical-abstract"
HL_DEL = ROOT / "paper" / "delivery" / "attachments" / "highlights"
TEXBIN = Path(r"C:\Users\wang\AppData\Local\Programs\MiKTeX\miktex\bin\x64")

BULLETS = [
    "Train-once export from MCU flash/SRAM budgets to Pi ONNX",
    "ESP-DL-legal graph split with identity-skip residual fix",
    "Measured Pi/ESP32 F1--latency, stage share, and 5 V mJ/sentence",
    "Memory map and reproducible UART benchmarking protocol",
]

GA_W, GA_H = 2000, 800  # Elsevier 5:2, above 1328x531


def write_highlights_text() -> None:
    for b in BULLETS:
        assert len(b) <= 85, (len(b), b)
        print(f"{len(b):2d} chars: {b}")

    md = '---\ntitle: "Highlights"\n---\n\n' + "\n".join(f"- {b}" for b in BULLETS) + "\n"
    txt = "\n".join(BULLETS) + "\n"
    for d in (SUB, HL_DEL):
        d.mkdir(parents=True, exist_ok=True)
        (d / "highlights.md").write_text(md, encoding="utf-8")
        (d / "highlights.txt").write_text(txt, encoding="utf-8")

    doc = Document()
    doc.add_heading("Highlights", level=1)
    for b in BULLETS:
        doc.add_paragraph(b, style="List Bullet")
    for d in (SUB, HL_DEL):
        doc.save(d / "highlights.docx")
    print("wrote highlights md/txt/docx")


def run_pdflatex(tex_name: str) -> None:
    env = os.environ.copy()
    env["PATH"] = str(TEXBIN) + os.pathsep + env.get("PATH", "")
    cmd = [
        str(TEXBIN / "pdflatex.exe"),
        "-interaction=nonstopmode",
        "-halt-on-error",
        tex_name,
    ]
    r = subprocess.run(cmd, cwd=SUB, env=env, capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout[-3000:])
        print(r.stderr[-2000:])
        raise SystemExit(f"pdflatex failed: {tex_name}")
    print(f"compiled {tex_name}")


def rasterize_ga() -> Path:
    env = os.environ.copy()
    env["PATH"] = str(TEXBIN) + os.pathsep + env.get("PATH", "")
    out_base = SUB / "graphical-abstract"
    cmd = [
        str(TEXBIN / "pdftoppm.exe"),
        "-png",
        "-r",
        "300",
        str(SUB / "graphical-abstract.pdf"),
        str(out_base),
    ]
    r = subprocess.run(cmd, env=env, capture_output=True, text=True)
    if r.returncode != 0:
        print(r.stdout, r.stderr)
        raise SystemExit("pdftoppm failed")

    cands = sorted(SUB.glob("graphical-abstract*.png"), key=lambda p: p.stat().st_mtime, reverse=True)
    if not cands:
        raise SystemExit("no GA png")
    target = SUB / "graphical-abstract.png"
    if cands[0].resolve() != target.resolve():
        cands[0].replace(target)
    for p in SUB.glob("graphical-abstract-*.png"):
        p.unlink(missing_ok=True)

    img = Image.open(target).convert("RGB")
    print(f"raw GA PNG: {img.size}")
    if img.size != (GA_W, GA_H):
        canvas = Image.new("RGB", (GA_W, GA_H), (255, 255, 255))
        scale = min(GA_W / img.width, GA_H / img.height)
        nw = max(1, int(img.width * scale))
        nh = max(1, int(img.height * scale))
        resized = img.resize((nw, nh), Image.Resampling.LANCZOS)
        canvas.paste(resized, ((GA_W - nw) // 2, (GA_H - nh) // 2))
        img = canvas
        print(f"padded GA PNG -> {img.size}")
    img.save(target, format="PNG", dpi=(300, 300))
    # Keep the LaTeX vector PDF (geometry already locks 5:2 page size).
    return target


def sync_all() -> None:
    GA_DEL.mkdir(parents=True, exist_ok=True)
    HL_DEL.mkdir(parents=True, exist_ok=True)

    shutil.copy2(SUB / "highlights.tex", HL_DEL / "highlights.tex")
    shutil.copy2(SUB / "highlights.pdf", HL_DEL / "highlights.pdf")
    shutil.copy2(SUB / "highlights.pdf", ROOT / "Highlights.pdf")

    shutil.copy2(SUB / "graphical-abstract.tex", GA_DEL / "graphical-abstract.tex")
    shutil.copy2(SUB / "graphical-abstract.png", GA_DEL / "graphical-abstract.png")
    shutil.copy2(SUB / "graphical-abstract.pdf", GA_DEL / "graphical-abstract.pdf")
    shutil.copy2(SUB / "graphical-abstract.pdf", ROOT / "GraphicalAbstract.pdf")

    print("OK synced all destinations")
    for p in [
        SUB / "highlights.md",
        SUB / "highlights.txt",
        SUB / "highlights.pdf",
        SUB / "highlights.docx",
        SUB / "highlights.tex",
        SUB / "graphical-abstract.png",
        SUB / "graphical-abstract.pdf",
        SUB / "graphical-abstract.tex",
        HL_DEL / "highlights.pdf",
        HL_DEL / "highlights.txt",
        GA_DEL / "graphical-abstract.png",
        GA_DEL / "graphical-abstract.pdf",
        ROOT / "Highlights.pdf",
        ROOT / "GraphicalAbstract.pdf",
    ]:
        print(f"  {p.relative_to(ROOT)}  {p.stat().st_size} bytes")


def main() -> None:
    write_highlights_text()
    run_pdflatex("highlights.tex")
    run_pdflatex("graphical-abstract.tex")
    rasterize_ga()
    sync_all()


if __name__ == "__main__":
    main()
